#!/usr/bin/env python3
"""Download the SharePoint L2 roster and write roster.json for the public board."""

from __future__ import annotations

import base64
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "roster.json"
TZ_NAME = "Asia/Kolkata"

SHARE_URLS = [
    "https://mckinsey-my.sharepoint.com/:x:/p/sophia_solomon_external/cQoJ_ECfjf9rS5LS-qnn9yLSEgUChleFsg8MpdMPWL_MJpJb0A",
    "https://mckinsey-my.sharepoint.com/:x:/r/personal/sophia_solomon_external_mckinsey_com/_layouts/15/Doc.aspx?sourcedoc=%7B9F40FC09-FF8D-4B6B-92D2-FAA9E7F722D2%7D&file=Eclipt%20OS%20L2%20Roster.xlsx&action=default&mobileredirect=true&share=cQoJ_ECfjf9rS5LS-qnn9yLSEgUChleFsg8MpdMPWL_MJpJb0A",
]

SHIFT_RULES = [
    (re.compile(r"week\s*end|on[-\s]?call", re.I), "weekend"),
    (re.compile(r"night|shift\s*3", re.I), "night"),
    (re.compile(r"evening|aft|shift\s*2", re.I), "evening"),
    (re.compile(r"morning|shift\s*1", re.I), "morning"),
]


def share_id(url: str) -> str:
    encoded = base64.b64encode(url.encode()).decode().rstrip("=")
    return "u!" + encoded.replace("+", "-").replace("/", "_")


def access_token(resource: str) -> str:
    env_key = "GRAPH_TOKEN" if "graph.microsoft.com" in resource else "SHAREPOINT_TOKEN"
    env = os.environ.get(env_key, "").strip()
    if env:
        return env
    try:
        raw = subprocess.check_output(
            [
                "az",
                "account",
                "get-access-token",
                "--resource",
                resource,
                "--query",
                "accessToken",
                "-o",
                "tsv",
            ],
            stderr=subprocess.DEVNULL,
            text=True,
        ).strip()
    except (subprocess.CalledProcessError, FileNotFoundError) as exc:
        if os.environ.get("GITHUB_ACTIONS") == "true":
            raise RuntimeError("missing-ci-token") from exc
        raise SystemExit(
            f"No token for {resource}. Sign in with `az login` or set {env_key}."
        ) from exc
    if not raw:
        raise SystemExit(f"Azure CLI returned an empty token for {resource}. Run `az login`.")
    return raw


def http_get(url: str, token: str, accept: str = "*/*") -> bytes:
    req = Request(url, headers={"Authorization": f"Bearer {token}", "Accept": accept})
    with urlopen(req, timeout=60) as resp:
        return resp.read()


def download_xlsx(dest: Path) -> None:
    errors: list[str] = []
    graph = access_token("https://graph.microsoft.com")

    for url in SHARE_URLS:
        for suffix in ("driveItem/content", "root/content"):
            endpoint = f"https://graph.microsoft.com/v1.0/shares/{share_id(url)}/{suffix}"
            try:
                data = http_get(endpoint, graph)
                if data[:2] == b"PK":
                    dest.write_bytes(data)
                    return
                errors.append(f"graph {suffix}: not xlsx")
            except HTTPError as exc:
                errors.append(f"graph {suffix}: {exc.code}")
            except URLError as exc:
                errors.append(f"graph {suffix}: {exc.reason}")

    for query_url in (
        "https://graph.microsoft.com/v1.0/me/drive/sharedWithMe?$top=50",
        "https://graph.microsoft.com/v1.0/me/drive/root/search(q='{}')".format(quote("Eclipt OS L2 Roster")),
    ):
        try:
            payload = json.loads(http_get(query_url, graph, accept="application/json"))
            for item in payload.get("value", []):
                name = str(item.get("name") or "")
                remote = item.get("remoteItem") or item
                if "Roster" not in name or not name.lower().endswith(".xlsx"):
                    continue
                parent = remote.get("parentReference") or {}
                drive_id = parent.get("driveId")
                item_id = remote.get("id")
                if not drive_id or not item_id:
                    continue
                data = http_get(
                    f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/content",
                    graph,
                )
                if data[:2] == b"PK":
                    dest.write_bytes(data)
                    return
        except HTTPError as exc:
            errors.append(f"graph search: {exc.code}")
        except URLError as exc:
            errors.append(f"graph search: {exc.reason}")

    try:
        spo = access_token("https://mckinsey-my.sharepoint.com")
    except SystemExit:
        spo = ""
    if spo:
        for url in SHARE_URLS:
            endpoint = (
                "https://mckinsey-my.sharepoint.com/_api/v2.0/shares/"
                f"{share_id(url)}/driveItem/content"
            )
            try:
                data = http_get(endpoint, spo)
                if data[:2] == b"PK":
                    dest.write_bytes(data)
                    return
                errors.append("sharepoint: not xlsx")
            except HTTPError as exc:
                errors.append(f"sharepoint: {exc.code}")
            except URLError as exc:
                errors.append(f"sharepoint: {exc.reason}")

    raise SystemExit("Could not download the SharePoint roster (" + "; ".join(errors) + ").")


def classify_shift(text: object) -> str:
    value = str(text or "").strip()
    if not value:
        return ""
    for pattern, label in SHIFT_RULES:
        if pattern.search(value):
            return label
    return ""


def to_iso(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 40000 < float(value) < 60000:
        base = date(1899, 12, 30) + timedelta(days=int(value))
        return base.isoformat()
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %B %Y", "%d %b %Y"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
    return ""


def split_names(text: object) -> list[str]:
    raw = str(text or "").strip()
    if not raw or classify_shift(raw) or re.match(r"^shift\b", raw, re.I):
        return []
    parts = re.split(r"\s*(?:,|/|&|\n| and )\s*", raw)
    names = []
    for part in parts:
        cleaned = re.sub(r"^@+", "", part).strip()
        if cleaned and not classify_shift(cleaned):
            names.append(cleaned)
    return names


def preferred_sheets(wb) -> list[str]:
    ranked = []
    for name in wb.sheetnames:
        score = 0
        if re.search(r"roster|shift|shirt", name, re.I):
            score += 2
        if re.search(r"final", name, re.I):
            score += 1
        if re.search(r"rule", name, re.I):
            score -= 2
        ranked.append((score, name))
    ranked.sort(reverse=True)
    return [name for score, name in ranked if score >= 0] or list(wb.sheetnames)


def parse_sheet(ws) -> tuple[dict[str, dict[str, str]], list[dict]]:
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    days: dict[str, dict[str, str]] = {}
    weekends: list[dict] = []
    date_row = -1
    date_cols: list[tuple[int, str]] = []

    for r, row in enumerate(rows[:25]):
        cols = []
        for c, cell in enumerate(row):
            iso = to_iso(cell)
            if iso:
                cols.append((c, iso))
        if len(cols) >= 3:
            date_row = r
            date_cols = cols
            break

    if date_row < 0:
        return days, weekends

    for row in rows[date_row + 1 :]:
        label = classify_shift(row[0] if row else "") or classify_shift(row[1] if len(row) > 1 else "")
        person = str(row[0] or "").strip() if row else ""
        for c, iso in date_cols:
            value = row[c] if c < len(row) else ""
            names = split_names(value)
            if label == "weekend" and names:
                weekends.append({"start": iso, "end": iso, "names": names})
            elif label in {"morning", "evening", "night"} and names:
                days.setdefault(iso, {})[label] = " · ".join(names)
            elif not label:
                maybe = classify_shift(value)
                if maybe == "weekend" and person:
                    weekends.append({"start": iso, "end": iso, "names": [person]})
                elif maybe in {"morning", "evening", "night"} and person:
                    days.setdefault(iso, {})[maybe] = person
    return days, weekends


def merge_weekends(items: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for item in items:
        key = "|".join(item["names"])
        current = grouped.get(key)
        if not current:
            grouped[key] = dict(item)
            continue
        current["start"] = min(current["start"], item["start"])
        current["end"] = max(current["end"], item["end"])
    return list(grouped.values())


def infer_pattern(days: dict[str, dict[str, str]]) -> dict[str, str]:
    counts = {key: defaultdict(int) for key in ("morning", "evening", "night")}
    for iso, slots in days.items():
        weekday = date.fromisoformat(iso).weekday()
        if weekday >= 5:
            continue
        for key in counts:
            if slots.get(key):
                counts[key][slots[key]] += 1
    return {
        key: max(values, key=values.get) if values else ""
        for key, values in counts.items()
    }


def parse_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    days: dict[str, dict[str, str]] = {}
    weekends: list[dict] = []
    used = []
    for name in preferred_sheets(wb):
        sheet_days, sheet_weekends = parse_sheet(wb[name])
        if sheet_days or sheet_weekends:
            used.append(name)
            for iso, slots in sheet_days.items():
                days.setdefault(iso, {}).update(slots)
            weekends.extend(sheet_weekends)
    if not days and not weekends:
        raise SystemExit(
            f"Could not find dates and shift rows in sheets: {', '.join(wb.sheetnames)}"
        )
    now = datetime.now(timezone(timedelta(hours=5, minutes=30)))
    return {
        "timezone": TZ_NAME,
        "source": f"SharePoint roster · synced {now.strftime('%d %b %Y, %H:%M')} IST",
        "syncedAt": now.isoformat(),
        "sheets": used,
        "shifts": [
            {"id": "morning", "label": "Morning", "start": "06:00", "end": "15:00"},
            {"id": "evening", "label": "Evening", "start": "15:00", "end": "22:00"},
            {"id": "night", "label": "Night", "start": "22:00", "end": "06:00"},
        ],
        "weekdayPattern": infer_pattern(days),
        "days": dict(sorted(days.items())),
        "weekends": sorted(merge_weekends(weekends), key=lambda item: item["start"]),
    }


def main() -> None:
    with TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "roster.xlsx"
        try:
            download_xlsx(xlsx)
        except RuntimeError as exc:
            if str(exc) == "missing-ci-token":
                print("Skipping SharePoint download: no GRAPH_TOKEN or Azure login in this environment.")
                return
            raise
        roster = parse_workbook(xlsx)
    previous = {}
    if OUTPUT.exists():
        previous = json.loads(OUTPUT.read_text())
    comparable_old = {k: previous.get(k) for k in ("days", "weekends", "weekdayPattern")}
    comparable_new = {k: roster.get(k) for k in ("days", "weekends", "weekdayPattern")}
    OUTPUT.write_text(json.dumps(roster, indent=2) + "\n")
    changed = comparable_old != comparable_new
    print(
        f"Wrote {OUTPUT.name}: {len(roster['days'])} days, "
        f"{len(roster['weekends'])} weekend rows, changed={changed}"
    )


if __name__ == "__main__":
    main()
